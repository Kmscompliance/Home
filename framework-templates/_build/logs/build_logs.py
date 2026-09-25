#!/usr/bin/env python3
"""Build the KMS v2.0 supporting registers in KMS house style."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xl import *

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'logsrc')
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'out')
BASE_APTOS = f'{SRC}/Risk_Assessment_Log_v1.0.xlsx'
YN = ['Yes', 'No']
YNNA = ['Yes', 'No', 'Not applicable']

ACTIVITIES = [
    ('Credit broking', 'your firm carries on credit broking.'),
    ('Debt adjusting and debt counselling', 'your firm carries on debt adjusting and/or debt counselling.'),
    ('Debt collecting and debt administration', 'your firm carries on debt collecting and/or debt administration.'),
    ('Credit information services and credit references', 'your firm provides credit information services and/or credit references.'),
    ('Peer-to-peer lending platforms', 'your firm operates an electronic system in relation to lending.'),
    ('Consumer credit lending', 'your firm enters into regulated credit agreements as lender.'),
    ('Exercising a lender’s rights and duties', 'your firm exercises the lender’s rights and duties under regulated credit agreements.'),
    ('Consumer hire', 'your firm hires out goods under regulated consumer hire agreements.'),
    ('Insurance distribution', 'your firm carries on insurance distribution.'),
    ('Mortgage intermediation', 'your firm advises on and/or arranges regulated mortgage contracts.'),
    ('Payment services and e-money', 'your firm provides payment services or issues electronic money.'),
    ('Investment advice and arranging', 'your firm advises on and/or arranges deals in investments.'),
]
ACT_NAMES = ['All activities'] + [a for a, _ in ACTIVITIES]


def keep_book(path, font, keep):
    """Open an existing KMS workbook, keep named sheets, remove the rest."""
    b = Book.__new__(Book)
    from openpyxl import load_workbook
    b.wb = load_workbook(path)
    b.font = font
    for ws in list(b.wb.worksheets):
        if ws.title not in keep:
            b.wb.remove(ws)
    b.lists = None
    b.list_cols = {}
    return b


# =====================================================================
# 1. Compliance Breach Log (updated)
# =====================================================================
def breach_log():
    b = Book(f'{SRC}/Compliance_Breach_Log_v1.0.xlsx', 'Calibri')
    b.guide('How to use', 'Compliance Breach Log – how to use', STANDARD_GUIDE + [
        ('h', 'What to record'),
        ('p', 'Record every breach of FCA rules, the Principles (including the Consumer Duty), the law (for example data protection or anti-money laundering law), the Conduct Rules, or our own policies – however minor, and whether or not customers were affected. Record near misses too where they reveal a weakness.'),
        ('h', 'Notifying the FCA'),
        ('p', 'For every breach, decide and record whether it must be notified to the FCA. Under Principle 11 and SUP 15.3 we must tell the FCA promptly about anything it would reasonably expect notice of, including significant rule breaches, matters that could affect our ability to meet the Threshold Conditions, and significant events such as fraud or serious operational failures. SM&CR firms must notify Conduct Rules breaches that lead to disciplinary action (SUP 15.11). If unsure, notify – and record your reasoning either way.'),
        ('watch', 'The FCA’s operational incident and third-party reporting rules (PS26/2) apply from 18 March 2027 and introduce standard incident reports. Review this log and your notification process before that date.'),
        ('h', 'Personal data breaches'),
        ('p', 'Personal data breaches are recorded in the separate Data Breach Log, which covers ICO notification within 72 hours. Where a data breach is also a regulatory breach (for example a significant cyber incident), record it in both and cross-refer.'),
        ('h', 'Putting it right'),
        ('p', 'Every breach must have remedial action with an owner and a deadline. Larger actions are tracked in the separate Corrective Action Log – record the reference here. Close a breach only when the root cause has been fixed and any affected customers have been put right.'),
    ])
    b.add_list('type', ['FCA rule / Principle', 'Consumer Duty outcome', 'Financial crime (AML / sanctions / fraud)', 'Data protection', 'Conduct Rules (SM&CR)', 'Financial promotion', 'Complaints handling (DISP)', 'Client money / safeguarding', 'Regulatory reporting / notification', 'Internal policy or procedure', 'Other'])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('fca', ['Yes – notified', 'Yes – to be notified', 'No – not required (reason recorded)', 'Under review'])
    b.add_list('sev', ['Red – significant / customer harm', 'Amber – moderate', 'Green – minor'])
    b.add_list('status', ['Open', 'Remediation in progress', 'Closed'])
    cols = ['No.', 'Date breach occurred', 'Date identified', 'How identified (and by whom)', 'Type of breach', 'Rule, Principle or policy breached', 'Description of the breach', 'Root cause', 'Severity',
            'No. of customers affected', 'Customer harm? (Yes/No)', 'Nature of harm and redress (amount, date paid)', 'Affected customers contacted?',
            'FCA notification decision', 'Date FCA notified / reason not notified', 'Personal data involved? (If yes, also record in Data Breach Log)',
            'Remedial actions taken', 'Corrective Action Log ref.', 'Owner', 'Target date', 'Status', 'Date closed', 'Closure signed off by']
    widths = [7, 12, 12, 22, 20, 22, 36, 28, 16, 11, 11, 28, 13, 20, 22, 18, 34, 13, 16, 12, 14, 12, 16]
    ws = b.sheet('Compliance Breach Log', widths, 'Compliance Breach Log')
    r = b.note(ws, 3, 'Record every breach and near miss. Every breach needs a recorded FCA notification decision (Principle 11 / SUP 15). See the “How to use” sheet.', len(cols))
    r = b.header(ws, r, cols, height=60, groups=[('Breach details', 9), ('Impact on customers', 4), ('Notifications', 3), ('Measures taken', 7)])
    end = b.rows(ws, r, 60, len(cols), height=32, first_ref='CB-')
    for key, col in [('type', 5), ('sev', 9), ('yn', 11), ('ynna', 13), ('fca', 14), ('yn', 16), ('status', 21)]:
        b.validate(ws, key, col, r, end - 1)
    b.rag(ws, f'I{r}:I{end - 1}')
    b.save(f'{OUT}/Compliance_Breach_Log_v2.0_Template.xlsx')


# =====================================================================
# 2. Customer Due Diligence Checklist (updated) – checklists, CDD Register, Sanctions log
# =====================================================================
CDD_ACTIVITY_ROWS = {
    'Credit broking': [('Application accuracy', 'Information submitted to lenders matches the verified evidence; nothing has been altered, added or “improved”.'), ('Joint applicants / guarantors', 'Every joint applicant and guarantor identified and verified before submission.')],
    'Debt adjusting and debt counselling': [('Authority to act', 'Any third party acting for the customer has written authority; authority checked.'), ('Payments into the plan', 'Payments come from the customer or an explained third party (client money protection).')],
    'Debt collecting and debt administration': [('Identity before discussing the debt', 'Identity confirmed through security checks before any account details were discussed.'), ('Third-party or lump-sum payment', 'Payer identified and payment explained where above [Insert threshold] or from a third party.')],
    'Credit information services and credit references': [('Multi-factor verification', 'Identity verified using at least two factors before any credit information was disclosed or changed.'), ('Change of contact details', 'Extra checks completed where the request came from a new device, address or contact detail.')],
    'Peer-to-peer lending platforms': [('Investor source of funds', 'Source of funds evidenced for deposits above [Insert threshold] in [Insert period].'), ('Withdrawal account', 'Withdrawal account verified as in the investor’s own name.'), ('Borrower beneficial owners', 'Beneficial owners of business borrowers identified and verified.')],
    'Consumer credit lending': [('Timing', 'CDD completed before the agreement was entered into (or delayed verification approved and completed before funds released).'), ('Introduced business', 'Where introduced by a broker or merchant, our own CDD completed (not reliant on the introducer unless a reg. 39 reliance agreement is in place).')],
    'Exercising a lender’s rights and duties': [('Inherited CDD', 'Seller’s CDD reviewed; missing or unreliable records refreshed at first contact.'), ('Lump-sum settlement', 'Source of funds checked for settlements above [Insert threshold].')],
    'Consumer hire': [('Release of goods', 'Photographic ID checked when goods were released or collected.'), ('Delivery address', 'Delivery address verified for goods above [Insert value].')],
    'Insurance distribution': [('Premium payer', 'Premium payer identified where not the policyholder.'), ('Life / investment-related insurance', 'Full CDD completed and beneficiaries identified (MLR business only). [Delete if you only distribute general insurance.]')],
    'Mortgage intermediation': [('Deposit', 'Source of deposit evidenced; gifted deposit donor identified and gift letter obtained.'), ('Income', 'Income evidence consistent with the application.')],
    'Payment services and e-money': [('Transfer of funds information', 'Payer and payee information collected and checked as the Transfer of Funds provisions require.'), ('Expected activity', 'Expected transaction volumes and destinations recorded for monitoring.')],
    'Investment advice and arranging': [('Source of funds', 'Source of funds evidenced for this investment.'), ('Source of wealth', 'Source of wealth established for high-risk clients and all PEPs.')],
}


def cdd_workbook():
    b = Book(f'{SRC}/Customer_Due_Dilligence_Checklist_v1.0.xlsx', 'Aptos Narrow')
    b.guide('How to use', 'Customer Due Diligence – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', '1. CDD Checklist – complete one copy for each customer at onboarding (copy the sheet, or save a copy of the workbook in the customer file).'),
        ('p', '2. EDD Checklist – complete in addition to the CDD Checklist whenever an enhanced due diligence trigger applies (see the Customer Due Diligence Policy, section 7).'),
        ('p', '3. CDD Register – one line per customer, giving a single view of every customer’s verification, risk rating, screening, review and deletion dates. This is the “CDD Register” referred to in the Anti-Money Laundering and Customer Due Diligence policies.'),
        ('p', '4. Sanctions Screening Log – one line per screening check, with potential matches and decisions. This is the “Sanctions Screening Log” referred to in the Anti-Money Laundering Policy.'),
        ('h', 'Firms outside the Money Laundering Regulations'),
        ('p', 'If your firm is outside the MLRs (see section 2 of your Anti-Money Laundering Policy), you must still verify identity and screen for sanctions, but you may simplify the EDD Checklist to match your Customer Due Diligence Policy. Do not delete the CDD Register or Sanctions Screening Log – they evidence your fraud and sanctions controls.'),
        ('h', 'Retention'),
        ('p', 'Keep CDD records for five years after the end of the business relationship or occasional transaction. Firms within the MLRs must then delete personal data unless the law, legal proceedings or the individual’s consent allow longer retention. The CDD Register calculates the deletion date for you.'),
    ])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('method', ['Electronic verification', 'Documentary – seen in person', 'Documentary – seen remotely (app / video)', 'Alternative evidence (approved)', 'Reliance (reg. 39)'])
    b.add_list('risk', ['Low', 'Standard', 'High'])
    b.add_list('level', ['Simplified (SDD)', 'Standard', 'Enhanced (EDD)'])
    b.add_list('pep', ['Not a PEP', 'Domestic (UK) PEP', 'Foreign PEP', 'Family member of a PEP', 'Known close associate of a PEP'])
    b.add_list('match', ['No match', 'Potential match – under review', 'False positive – cleared', 'Confirmed match – escalated'])
    b.add_list('ctype', ['Individual', 'Sole trader / partnership', 'Company / LLP', 'Trust / charity / other'])
    b.add_list('act', ACT_NAMES)
    b.add_list('role', ['Customer', 'Joint customer / guarantor', 'Beneficial owner', 'Person acting on behalf', 'Third-party payer', 'Payee'])

    # ---- CDD Checklist
    cols = ['Section', 'Checklist item', 'What must be done', 'Completed (Yes / No / N/A)', 'Date completed', 'Staff initials', 'Evidence / notes / follow-up required']
    widths = [22, 30, 60, 13, 13, 10, 45]
    ws = b.sheet('CDD Checklist', widths, 'Customer Due Diligence Checklist', 'Customer name: [Insert]   |   Customer reference: [Insert]   |   Firm: [Insert firm legal name]')
    r = b.note(ws, 3, 'Complete for every customer before an application is submitted, an agreement is entered into, or a transaction is carried out. Nothing proceeds until every applicable line is complete. If CDD cannot be completed, stop and consider an internal suspicion report (policy section 8).', len(cols))
    r = b.header(ws, r, cols)
    first = r
    universal = [
        ('1. Identification', 'Customer information collected', 'Full name, date of birth, residential address and contact details recorded. For business customers: registered name and number, registered office, principal place of business and directors/partners.'),
        ('1. Identification', 'Beneficial owners (business customers)', 'Individuals owning or controlling more than 25%, or otherwise exercising control, identified and verified; ownership structure understood.'),
        ('1. Identification', 'Persons acting on behalf', 'Anyone acting for the customer identified and their authority checked (e.g. power of attorney, company authority).'),
        ('2. Verification', 'Identity verified', 'Verified using a reliable, independent source. Record the method (electronic / documentary / alternative) and the provider or documents used.'),
        ('2. Verification', 'Address verified', 'Verified electronically or using a document dated within the last three months (see policy Appendix A).'),
        ('2. Verification', 'Remote customer checks', 'Where the customer was not seen in person: liveness / video check or additional verification completed. [Delete if you only deal face-to-face.]'),
        ('2. Verification', 'Documents checked', 'Documents appear genuine and unaltered; photo is a true likeness; details consistent with the application.'),
        ('3. Understanding', 'Purpose and intended nature', 'Why the customer is using our service, and expected activity or transaction size, recorded.'),
        ('3. Understanding', 'Third-party payer', 'Any person paying on the customer’s behalf identified, relationship established and explanation recorded.'),
        ('4. Screening', 'Sanctions screening', 'Customer (and beneficial owners / third-party payers) screened against the UK sanctions list; result recorded in the Sanctions Screening Log.'),
        ('4. Screening', 'PEP screening', 'PEP status checked and recorded (not a PEP / domestic / foreign / family member / close associate).'),
        ('4. Screening', 'High-risk third country', 'Checked for links to a jurisdiction subject to a FATF “call for action” (mandatory EDD) or under increased FATF monitoring (risk factor).'),
        ('4. Screening', 'Adverse information and fraud indicators', 'Any warning signs (inconsistent information, pressure from others, unusual payment) considered and recorded.'),
        ('5. Risk decision', 'Customer risk rating', 'Low / standard / high rating assigned using the Business-Wide Risk Assessment factors, with reasons.'),
        ('5. Risk decision', 'Level of due diligence', 'SDD (with documented low-risk rationale – MLR firms only), standard CDD, or EDD (complete the EDD Checklist).'),
        ('5. Risk decision', 'Vulnerability noted', 'Any characteristics of vulnerability noted and support recorded in line with the Customer Vulnerability Policy (record needs, not diagnoses).'),
        ('6. Records', 'Privacy notice provided', 'Customer given our privacy notice (and, where relevant, the Credit Reference Agency Information Notice). Note: CDD is carried out to meet a legal obligation or for fraud prevention – do not rely on consent.'),
        ('6. Records', 'Customer Due Diligence Register updated', 'Customer entered in the Customer Due Diligence Register with risk rating, screening date and next review date.'),
        ('6. Records', 'Records stored securely', 'Evidence stored securely with retention date (five years after the relationship ends).'),
    ]
    for sec, item, what in universal:
        r = b.data_row(ws, r, [sec, item, what], len(cols), height=48)
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        for item, what in CDD_ACTIVITY_ROWS[act]:
            r = b.data_row(ws, r, ['7. Activity-specific', item, what], len(cols), height=40)
    r = b.data_row(ws, r, ['Sign-off', 'CDD complete and approved', 'All applicable items complete. Checked by: [Insert name]   Date: [Insert date]'], len(cols), height=36, bold=True)
    b.validate(ws, 'ynna', 4, first, r - 1)

    # ---- EDD Checklist
    ws = b.sheet('EDD Checklist', widths, 'Enhanced Due Diligence Checklist', 'Customer name: [Insert]   |   Customer reference: [Insert]   |   Firm: [Insert firm legal name]')
    r = b.note(ws, 3, 'Complete in addition to the CDD Checklist whenever an EDD trigger applies. For firms within the MLRs, senior management approval is mandatory before establishing or continuing a relationship with a PEP.', len(cols))
    r = b.header(ws, r, cols)
    first = r
    edd = [
        ('1. Trigger', 'Reason for EDD', 'Record the trigger(s): PEP / family member / close associate; FATF “call for action” jurisdiction; unusually complex or unusually large transaction or unusual pattern; no apparent economic or legal purpose; high risk rating; fraud indicators.'),
        ('2. Measures', 'Additional verification', 'Identity and address verified from additional independent sources.'),
        ('2. Measures', 'Source of funds', 'Evidence obtained of where the money for this transaction comes from (e.g. bank statements, payslips, completion statement, solicitor’s letter).'),
        ('2. Measures', 'Source of wealth', 'Where appropriate (always for foreign PEPs and high-risk investment clients), how the customer acquired their overall wealth established and evidenced.'),
        ('2. Measures', 'Purpose of transaction', 'Reason for the transaction explored and recorded, and found to make sense for this customer.'),
        ('2. Measures', 'Adverse media search', 'Search for negative news or public information completed and recorded.'),
        ('2. Measures', 'Enhanced screening', 'Further PEP / sanctions screening completed where the first result was unclear.'),
        ('3. Approval', 'MLRO consulted', 'MLRO (or nominated person) reviewed the case and recorded their view.'),
        ('3. Approval', 'Senior management approval', 'Approval given by [Insert name / role] before proceeding (mandatory for PEPs in MLR firms). Record decision and reasons.'),
        ('4. Ongoing', 'Enhanced monitoring', 'Closer monitoring set up and next review date recorded in the Customer Due Diligence Register (at least [Insert frequency]).'),
        ('4. Ongoing', 'Records', 'All EDD evidence and decisions stored securely; Customer Due Diligence Register updated.'),
    ]
    for sec, item, what in edd:
        r = b.data_row(ws, r, [sec, item, what], len(cols), height=48)
    r = b.data_row(ws, r, ['Decision', 'Proceed / do not proceed', 'Decision: [Insert: proceed / decline / exit]   Approved by: [Insert name]   Date: [Insert date]'], len(cols), height=36, bold=True)
    b.validate(ws, 'ynna', 4, first, r - 1)

    # ---- CDD Register
    cols = ['Customer ref.', 'Customer name', 'Customer type', 'Regulated activity / product', 'Date onboarded', 'Verification method', 'Date verified', 'Verified by', 'Customer risk rating', 'Level of due diligence', 'PEP status', 'High-risk country link? (details)', 'Third-party payer? (details)', 'Date of last sanctions screen', 'EDD approved by (and date)', 'Next review date', 'Date relationship ended', 'Delete records on (5 years after end)', 'Notes']
    widths = [12, 22, 16, 22, 12, 22, 12, 12, 11, 15, 18, 18, 18, 13, 18, 12, 13, 14, 30]
    ws = b.sheet('CDD Register', widths, 'Customer Due Diligence Register')
    r = b.note(ws, 3, 'One line per customer. Review dates: high-risk customers at least [Insert frequency]; others [Insert trigger/frequency]. The deletion date is calculated automatically five years after the relationship ends.', len(cols))
    r = b.header(ws, r, cols)
    end = b.rows(ws, r, 150, len(cols), height=20, formulas={18: '=IF(Q{r}="","",EDATE(Q{r},60))'})
    for key, col in [('ctype', 3), ('act', 4), ('method', 6), ('risk', 9), ('level', 10), ('pep', 11)]:
        b.validate(ws, key, col, r, end - 1)
    for col in (5, 7, 14, 16, 17, 18):
        for rr in range(r, end):
            ws.cell(rr, col).number_format = 'DD/MM/YYYY'

    # ---- Sanctions Screening Log
    cols = ['Screen ref.', 'Date screened', 'Name screened', 'Customer ref.', 'Role', 'Screening tool / list used', 'Result', 'Reviewed by', 'Decision and reasons', 'Account / transaction frozen?', 'Reported to OFSI? (date and reference)', 'MLRO informed? (date)']
    widths = [11, 12, 24, 12, 18, 22, 22, 14, 36, 12, 22, 14]
    ws = b.sheet('Sanctions Screening Log', widths, 'Sanctions Screening Log')
    r = b.note(ws, 3, 'Record every screening (at onboarding and each re-screen). A potential match must be escalated to the MLRO immediately, and nothing further done on the account or transaction until the MLRO has decided. As an FCA-authorised firm we must report to OFSI as soon as practicable if we know or reasonably suspect a person is designated or has breached sanctions.', len(cols))
    r = b.header(ws, r, cols)
    end = b.rows(ws, r, 150, len(cols), height=20, first_ref='SS-')
    for key, col in [('role', 5), ('match', 7), ('ynna', 10)]:
        b.validate(ws, key, col, r, end - 1)
    b.save(f'{OUT}/Customer_Due_Diligence_Checklist_v2.0_Template.xlsx')


# =====================================================================
# 3. Risk Assessment Log (updated)
# =====================================================================
UNIVERSAL_RISKS = [
    ('Conduct and customer harm', 'Poor customer outcomes', 'Customers receive products, services or support that do not meet their needs, leading to financial loss or distress (Consumer Duty).', 'Target market and product approval; outcome MI; file reviews; complaints root cause analysis.'),
    ('Conduct and customer harm', 'Vulnerable customers not identified or supported', 'Staff fail to recognise or respond to characteristics of vulnerability, so these customers receive worse outcomes.', 'Vulnerability Policy; training; recording and flags; outcome comparison MI.'),
    ('Regulatory and legal', 'Non-compliant financial promotions', 'A promotion (including social media or third-party content) is misleading or omits required information.', 'Pre-approval by competent approver; Financial Promotions Log; periodic review of live promotions.'),
    ('Regulatory and legal', 'Complaints mishandled', 'Complaints are not recognised, are handled late or unfairly, or root causes are not fixed.', 'Complaints Procedure; Complaints Register with deadline tracking; quarterly review.'),
    ('Regulatory and legal', 'Regulatory change missed', 'New or changed rules are not identified and implemented on time.', 'Regulatory updates from [Insert source]; Compliance Universe Register; external compliance support.'),
    ('Financial crime', 'Firm used for money laundering or fraud', 'Criminals use our services to launder money or commit fraud (application, identity or payment fraud).', 'CDD and EDD; sanctions screening; staff training; suspicion reporting to MLRO.'),
    ('Operational', 'Key person dependency', 'The loss or absence of [Insert name / role] disrupts regulated activity or oversight.', 'Documented procedures; cover arrangements; business continuity plan.'),
    ('Operational', 'Outsourcing / supplier failure', 'A critical supplier (e.g. IT, CRM, lender or insurer platform) fails or performs poorly.', 'Supplier due diligence; contracts; monitoring; exit plans.'),
    ('Technology, cyber and data', 'Cyber attack or data breach', 'Customer personal data is lost, stolen or disclosed, or systems are unavailable.', 'Multi-factor authentication; encryption; patching; back-ups; training; breach procedure.'),
    ('Financial and prudential', 'Insufficient financial resources', 'Income falls or costs (including redress) rise so that we cannot meet the Threshold Conditions or wind down in an orderly way.', 'Cash-flow forecasting; minimum buffer; stress testing; wind-down plan.'),
    ('Strategic and business model', 'Business model not viable', 'Market, competition or regulatory change makes the business model unsustainable.', 'Business plan review; board oversight; diversification of income.'),
    ('Reputational', 'Loss of trust', 'Poor outcomes, complaints or media coverage damage our reputation with customers, partners or the FCA.', 'Consumer Duty monitoring; complaints handling; open relationship with the FCA.'),
]
ACTIVITY_RISKS = {
    'Credit broking': [('Conduct and customer harm', 'Inadequate explanations / pressure selling', 'Customers enter credit agreements they do not understand or that are unsuitable.', 'Scripts; CONC 4.2 training; file reviews.'), ('Regulatory and legal', 'Commission disclosure failures', 'Commission arrangements are not disclosed adequately, leading to complaints and redress.', 'Standard disclosures; lender panel review; monitoring.')],
    'Debt adjusting and debt counselling': [('Conduct and customer harm', 'Unsuitable debt advice', 'Customers are put into unsuitable or unsustainable solutions.', 'Qualified advisers; file reviews; free-advice signposting.'), ('Financial and prudential', 'Client money shortfall', 'Client money is lost, misused or not distributed promptly (CASS 11).', 'Daily reconciliations; segregation; audit.')],
    'Debt collecting and debt administration': [('Conduct and customer harm', 'Unfair collection practices', 'Customers in difficulty are treated without forbearance or pressured (CONC 7).', 'Call monitoring; forbearance procedures; training.'), ('Regulatory and legal', 'Pursuing wrong or disputed debts', 'Collection continues on disputed or statute-barred debts, or the wrong person.', 'ID checks; dispute suspension; data validation.')],
    'Credit information services and credit references': [('Conduct and customer harm', 'Inaccurate credit data', 'Inaccurate data causes customers to be refused credit or charged more.', 'Data quality checks; prompt dispute handling.'), ('Technology, cyber and data', 'Large-scale data breach', 'Large volumes of sensitive credit data are exposed.', 'Enhanced security; access controls; DPIAs.')],
    'Peer-to-peer lending platforms': [('Conduct and customer harm', 'Investors misunderstand risk', 'Investors do not understand the risk of loss or illiquidity.', 'COBS 4.12A warnings; appropriateness; cooling-off.'), ('Financial and prudential', 'Platform failure', 'The platform fails and loans cannot be administered.', 'IPRU-INV 12 capital; back-up servicer; wind-down plan.')],
    'Consumer credit lending': [('Conduct and customer harm', 'Irresponsible lending', 'Credit is granted that customers cannot afford (CONC 5).', 'Affordability model; underwriting mandates; outcome MI.'), ('Financial and prudential', 'Credit losses / funding', 'Defaults exceed expectations or funding is withdrawn.', 'Credit risk appetite; funding diversification; stress testing.')],
    'Exercising a lender’s rights and duties': [('Regulatory and legal', 'Inherited non-compliance', 'Acquired agreements have errors that make them unenforceable or require redress.', 'Pre-purchase due diligence; data validation.'), ('Operational', 'Servicer failure', 'Third-party servicer fails or treats customers unfairly.', 'Oversight MI; contractual rights; back-up arrangements.')],
    'Consumer hire': [('Conduct and customer harm', 'Unfair end-of-hire charges', 'Charges are unclear, unevidenced or disproportionate.', 'Clear pre-contract information; evidence-based charging; review.'), ('Financial and prudential', 'Residual value risk', 'Returned goods are worth less than expected.', 'Residual value assumptions reviewed; provisioning.')],
    'Insurance distribution': [('Conduct and customer harm', 'Poor value products / add-ons', 'Customers buy products that offer poor value (PROD 4).', 'Manufacturer fair value information; remuneration review.'), ('Financial and prudential', 'PII / capital shortfall', 'Professional indemnity insurance or capital falls below MIPRU requirements.', 'Annual PII review; capital monitoring.')],
    'Mortgage intermediation': [('Conduct and customer harm', 'Unsuitable mortgage advice', 'Advice is unsuitable, including debt consolidation or long terms (MCOB 4).', 'File reviews; qualified advisers; supervision.'), ('Financial crime', 'Mortgage fraud', 'False income or deposit information is submitted to lenders.', 'Evidence checks; staff training; reporting.')],
    'Payment services and e-money': [('Financial and prudential', 'Safeguarding failure', 'Relevant funds are not segregated or reconciled (CASS 15).', 'Daily reconciliations; resolution pack; audit.'), ('Technology, cyber and data', 'Major operational or security incident', 'Payment services are disrupted or compromised.', 'Security risk assessment; incident reporting; impact tolerances.')],
    'Investment advice and arranging': [('Conduct and customer harm', 'Unsuitable advice', 'Advice is unsuitable for clients’ needs or capacity for loss (COBS 9 / 9A).', 'Suitability reviews; file checks; CPD.'), ('Conduct and customer harm', 'Ongoing service not delivered', 'Clients pay for ongoing advice they do not receive.', 'Service tracking; annual review evidence.')],
}


def risk_log():
    b = Book(f'{SRC}/Risk_Assessment_Log_v1.0.xlsx', 'Aptos Narrow')
    b.guide('How to use', 'Risk Assessment Log – how to use', STANDARD_GUIDE + [
        ('h', 'What this log is'),
        ('p', 'This is the Risk Assessment Log referred to in the Risk Management Framework. It records the risks your business model creates for customers, the market and the firm, how you control them, and whether they are within your risk appetite. (The AML Business-Wide Risk Assessment is a separate register.)'),
        ('h', 'Scoring (Risk Management Framework, section 5.2)'),
        ('p', 'Score each risk for impact (1–5) and likelihood (1–5), first before controls (inherent) and then after controls (residual). The log multiplies them. Residual scores of 15 or more are HIGH, 8 to 14 MEDIUM and 1 to 7 LOW.'),
        ('p', 'Impact: 1 Very low – negligible effect on customers or the firm; 2 Low – minor, easily corrected; 3 Medium – some customer detriment or regulatory breach; 4 High – significant customer harm, material loss or reportable breach; 5 Very high – widespread harm, threat to the firm’s viability or authorisation.'),
        ('p', 'Likelihood: 1 Rare (less than once in 5 years); 2 Unlikely (once in 2–5 years); 3 Possible (once a year); 4 Likely (several times a year); 5 Almost certain (monthly or more).'),
        ('h', 'Starter risks'),
        ('p', 'The log is pre-populated with common risks to help you start. They are examples, not a finished assessment: rewrite each in terms of your own business, delete any that do not apply, add your own, and score them. Activity-specific risks sit under green OPTION rows. The FCA expects a register that is clearly specific to your firm.'),
        ('p', 'Review the log at least [Insert frequency, e.g. quarterly] and report high and increasing risks to [Insert: the board / the directors / the business owner].'),
    ])
    b.add_list('cat', ['Conduct and customer harm', 'Regulatory and legal', 'Financial crime', 'Operational', 'Technology, cyber and data', 'Financial and prudential', 'Strategic and business model', 'Reputational'])
    b.add_list('score', [1, 2, 3, 4, 5])
    b.add_list('eff', ['Effective', 'Partly effective', 'Not effective', 'Not yet tested'])
    b.add_list('yn', YN)
    b.add_list('status', ['Open', 'Action in progress', 'Accepted within appetite', 'Closed'])
    b.add_list('act', ACT_NAMES)
    cols = ['Risk ID', 'Category', 'Activity', 'Risk summary', 'Risk description (cause → event → consequence)', 'Inherent impact (1–5)', 'Inherent likelihood (1–5)', 'Inherent score', 'Key controls', 'Control effectiveness', 'Residual impact (1–5)', 'Residual likelihood (1–5)', 'Residual score', 'Residual rating', 'Within appetite?', 'Risk owner', 'Further actions', 'Action due', 'Status', 'Last reviewed', 'Next review']
    widths = [9, 20, 18, 26, 44, 10, 10, 9, 38, 14, 10, 10, 9, 10, 10, 15, 30, 11, 14, 11, 11]
    ws = b.sheet('Risk Assessment', widths, 'Risk Assessment Log')
    r = b.note(ws, 3, 'Starter risks are examples – rewrite, score and add to them. Scores and ratings calculate automatically. Delete OPTION rows for activities you do not carry on.', len(cols))
    r = b.header(ws, r, cols, height=48, groups=[('Risk identification', 5), ('Inherent risk', 3), ('Controls', 2), ('Residual risk', 5), ('Ownership and action', 6)])
    first = r
    fx = {8: '=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})', 13: '=IF(OR(K{r}="",L{r}=""),"",K{r}*L{r})', 14: '=IF(M{r}="","",IF(M{r}>=15,"HIGH",IF(M{r}>=8,"MEDIUM","LOW")))'}
    n = 1
    def put(r, cat, act, summ, desc, ctrl):
        nonlocal n
        r2 = b.data_row(ws, r, [f'R{n:02d}', cat, act, summ, desc, None, None, None, ctrl, None, None, None, None, None, None, '[Insert owner]'], len(cols), height=60)
        for col, t in fx.items():
            ws.cell(r, col, t.format(r=r))
        n += 1
        return r2
    for cat, summ, desc, ctrl in UNIVERSAL_RISKS:
        r = put(r, cat, 'All activities', summ, desc, ctrl)
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        for cat, summ, desc, ctrl in ACTIVITY_RISKS[act]:
            r = put(r, cat, act, summ, desc, ctrl)
    blank_start = r
    for i in range(15):
        r = b.data_row(ws, r, [f'R{n:02d}'], len(cols), height=30)
        for col, t in fx.items():
            ws.cell(r - 1, col, t.format(r=r - 1))
        n += 1
    last = r - 1
    for key, col in [('cat', 2), ('act', 3), ('score', 6), ('score', 7), ('eff', 10), ('score', 11), ('score', 12), ('yn', 15), ('status', 19)]:
        b.validate(ws, key, col, first, last)
    for col in ('H', 'M'):
        rng = f'{col}{first}:{col}{last}'
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['15'], fill=PatternFill('solid', fgColor='F4B6B6')))
        ws.conditional_formatting.add(rng, CellIsRule(operator='between', formula=['8', '14'], fill=PatternFill('solid', fgColor='FFE699')))
        ws.conditional_formatting.add(rng, CellIsRule(operator='between', formula=['1', '7'], fill=PatternFill('solid', fgColor='C6E0B4')))
    rng = f'N{first}:N{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{first}="HIGH"'], fill=PatternFill('solid', fgColor='F4B6B6')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{first}="MEDIUM"'], fill=PatternFill('solid', fgColor='FFE699')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{first}="LOW"'], fill=PatternFill('solid', fgColor='C6E0B4')))
    b.save(f'{OUT}/Risk_Assessment_Log_v2.0_Template.xlsx')


# =====================================================================
# 4. Training Log (updated)
# =====================================================================
TRAINING_MODULES = [
    ('Induction – our business, permissions and customers', 'All staff', 'Before customer contact', 'Once'),
    ('FCA Principles, Consumer Duty and (SM&CR firms) Conduct Rules', 'All staff', 'Induction', 'Annually'),
    ('Anti-money laundering, sanctions and fraud', 'All staff', 'Induction', 'Annually'),
    ('Customer due diligence procedures', 'Onboarding / payments staff', 'Before onboarding customers', 'Annually'),
    ('Data protection and information security', 'All staff', 'Induction', 'Annually'),
    ('Customer vulnerability', 'Customer-facing staff', 'Before unsupervised customer contact', 'Annually'),
    ('Complaints recognition and handling', 'All staff (handlers: detailed)', 'Induction', 'Annually'),
    ('Financial promotions (including verbal and social media)', 'Sales and marketing staff; approver', 'Before creating / approving promotions', 'Annually'),
    ('Senior manager duties, Duty of Responsibility and governance', 'SMF holders', 'On appointment', 'Annually'),
]
TRAINING_ACT = {
    'Credit broking': ('CONC 3, 4.2, 4.4 and 4.5 – explanations, status and commission disclosure', 'No TC qualification; competence assessed internally'),
    'Debt adjusting and debt counselling': ('CONC 8 debt advice; Breathing Space; client money (CASS 11)', 'Recognised money advice qualification: [Insert]; CPD: [Insert] hours'),
    'Debt collecting and debt administration': ('CONC 7 forbearance (PS24/2); Standard Financial Statement; Breathing Space', 'No TC qualification; call monitoring before unsupervised work'),
    'Credit information services and credit references': ('CCA ss158–159 rights; disputes; identity verification', 'No TC qualification'),
    'Peer-to-peer lending platforms': ('COBS 18.12; COBS 4.12A; appropriateness (COBS 10); CASS 7', 'No TC qualification unless giving personal recommendations'),
    'Consumer credit lending': ('CONC 5 creditworthiness; CONC 4 and 6; CONC 7 arrears (PS24/2)', 'No TC qualification; underwriting mandates'),
    'Exercising a lender’s rights and duties': ('CONC 7; CCA notices and enforceability; servicer oversight', 'No TC qualification'),
    'Consumer hire': ('Hire agreements and end-of-hire charges; CONC 7', 'No TC qualification'),
    'Insurance distribution': ('ICOBS; PROD 4; claims support', 'TC: at least 15 hours CPD per year'),
    'Mortgage intermediation': ('MCOB 4 and 5A; affordability; later-life lending', 'TC: appropriate qualification (e.g. CeMAP); equity release qualification if applicable; CPD'),
    'Payment services and e-money': ('PSRs 2017; CASS 15 safeguarding; SCA; fraud and APP scams', 'No TC; competent staff and fit and proper management'),
    'Investment advice and arranging': ('COBS 9/9A; COBS 6; high-risk investments', 'TC: level 4 qualification; annual SPS; 35 hours CPD (21 structured)'),
}


def training_log():
    b = keep_book(f'{SRC}/Training_Log_v1.0.xlsx', 'Calibri', keep=['Employee 1'])
    ind = b.wb['Employee 1']
    ind.title = 'Individual Record (template)'
    for row in ind.iter_rows():
        for c in row:
            if c.value == "FIRM's NAME":
                c.value = '[Insert firm legal name]'
                c.fill = PH_FILL
    ind.cell(2, 1, 'Optional reflective record – copy this sheet for each person if you want to capture what they learned and how it changed their work. The Training Log is the mandatory record.').font = Font(name='Calibri', i=True, sz=10)
    b.guide('How to use', 'Training Log – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', '1. Training Matrix – the training each role must complete, and how often. This is the training matrix referred to in the Training and Competence Policy.'),
        ('p', '2. Training Record – one line for every training activity completed by anyone in the firm. The next-due date calculates automatically. This is the mandatory “Training Log”.'),
        ('p', '3. Competence Register – one line per person: whether they are assessed as competent, supervision, qualifications, CPD and (for SM&CR firms) fitness and propriety and certification dates.'),
        ('p', '4. CPD Log – for staff with CPD requirements (for example insurance distribution, mortgage and investment advisers). Delete it if no one in your firm has a CPD requirement.'),
        ('p', '5. Individual Record (template) – the original KMS reflective record, kept as an optional sheet.'),
        ('h', 'Evidence'),
        ('p', 'Keep certificates, assessment results and signed declarations. Retain records for [Insert period] after the person leaves, and for at least the minimum period TC requires where it applies (indefinitely for pension transfer specialists).'),
    ])
    b.add_list('type', ['Induction', 'Annual refresher', 'Role-specific', 'Regulatory change', 'Remedial / targeted', 'CPD / development'])
    b.add_list('method', ['E-learning', 'Face-to-face / webinar', 'Shadowing / coaching', 'Reading and declaration', 'External course / qualification'])
    b.add_list('result', ['Passed', 'Failed – retake required', 'Declaration signed', 'Not assessed'])
    b.add_list('sup', ['Fully supervised', 'Partly supervised', 'Assessed competent – unsupervised', 'Restricted'])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('struct', ['Structured', 'Unstructured'])
    # Matrix
    cols = ['Training module', 'Who must complete it', 'When first required', 'Refresher frequency', 'Delivered by / provider']
    widths = [58, 34, 30, 18, 30]
    ws = b.sheet('Training Matrix', widths, 'Training Matrix')
    r = b.note(ws, 3, 'Adapt the modules and roles to your firm. Activity-specific modules sit under green OPTION rows – delete those you do not need.', len(cols))
    r = b.header(ws, r, cols)
    for m in TRAINING_MODULES:
        r = b.data_row(ws, r, list(m) + ['[Insert]'], len(cols), height=30)
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        mod, req = TRAINING_ACT[act]
        r = b.data_row(ws, r, [f'{act}: {mod}', 'Staff carrying on this activity', 'Before working unsupervised', 'Annually', '[Insert]'], len(cols), height=30)
        r = b.data_row(ws, r, [f'{act}: qualification / CPD requirement – {req}', 'Staff carrying on this activity', 'As TC / firm standard requires', 'Annually', '[Insert]'], len(cols), height=30)
    # Record
    cols = ['Name', 'Role', 'Training module', 'Type', 'Date completed', 'Method / provider', 'Assessment result', 'Declaration of understanding signed?', 'Refresher due (months)', 'Next due', 'Evidence location', 'What was learned and how it will be applied']
    widths = [18, 16, 36, 16, 12, 20, 16, 13, 11, 12, 20, 40]
    ws = b.sheet('Training Log', widths, 'Training Log')
    r = b.note(ws, 3, 'One line per training activity completed. Enter the refresher period in months (e.g. 12) and the next-due date calculates automatically; overdue dates turn red.', len(cols))
    r = b.header(ws, r, cols)
    end = b.rows(ws, r, 200, len(cols), height=20, formulas={10: '=IF(OR(E{r}="",I{r}=""),"",EDATE(E{r},I{r}))'})
    for key, col in [('type', 4), ('method', 6), ('result', 7), ('yn', 8)]:
        b.validate(ws, key, col, r, end - 1)
    for rr in range(r, end):
        ws.cell(rr, 5).number_format = 'DD/MM/YYYY'; ws.cell(rr, 10).number_format = 'DD/MM/YYYY'
    ws.conditional_formatting.add(f'J{r}:J{end - 1}', FormulaRule(formula=[f'AND(J{r}<>"",J{r}<TODAY())'], fill=PatternFill('solid', fgColor='F4B6B6')))
    # Competence register
    cols = ['Name', 'Role', 'Regulated / customer-facing activities performed', 'Start date', 'Supervision status', 'Assessed competent on', 'Assessed by', 'Qualification required (TC or firm standard)', 'Qualification held and date', 'SPS expiry (investment advisers)', 'CPD hours required per year', 'CPD hours completed (from CPD Log)', 'SMF / certification function (SM&CR firms)', 'Fitness and propriety last assessed', 'Certificate issued / renewed', 'Next competence review']
    widths = [18, 16, 30, 11, 20, 12, 14, 24, 20, 12, 11, 12, 18, 13, 13, 13]
    ws = b.sheet('Competence Register', widths, 'Competence Register')
    r = b.note(ws, 3, 'One line per person. Nobody carries out a regulated or customer-facing activity unsupervised until the “Assessed competent on” date is completed. Delete the SM&CR columns if your firm is outside SM&CR (payment and e-money institutions), and the CPD/SPS columns if no one has a CPD requirement.', len(cols))
    r = b.header(ws, r, cols)
    end = b.rows(ws, r, 40, len(cols), height=24, formulas={12: "=IF(A{r}=\"\",\"\",SUMIFS('CPD Log'!E:E,'CPD Log'!A:A,A{r},'CPD Log'!B:B,\">=\"&DATE(YEAR(TODAY()),1,1)))"})
    b.validate(ws, 'sup', 5, r, end - 1)
    # CPD log
    cols = ['Name', 'Date', 'CPD activity', 'Structured / unstructured', 'Hours', 'Learning outcome', 'Evidence location']
    widths = [18, 12, 40, 16, 8, 40, 22]
    ws = b.sheet('CPD Log', widths, 'CPD Log')
    r = b.note(ws, 3, 'For staff with CPD requirements: insurance distribution (at least 15 hours a year), mortgage advisers ([Insert] hours), investment advisers (35 hours, of which 21 structured). The Competence Register totals hours for the current calendar year – adjust the formula if your CPD year differs (for example to match your SPS year).', len(cols))
    r = b.header(ws, r, cols)
    end = b.rows(ws, r, 200, len(cols), height=20)
    b.validate(ws, 'struct', 4, r, end - 1)
    b.wb.move_sheet('Individual Record (template)', offset=len(b.wb.sheetnames))
    b.save(f'{OUT}/Training_Log_v2.0_Template.xlsx')


# =====================================================================
# 5. Complaints Register (new)
# =====================================================================
def complaints_register():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Complaints Register – how to use', STANDARD_GUIDE + [
        ('h', 'What to record'),
        ('p', 'Record every complaint – any oral or written expression of dissatisfaction about our financial services alleging financial loss, material distress or material inconvenience – on the day it is received (or by the end of the next business day). The DISP time limits run from the date any member of staff, or anyone acting for us, first received it.'),
        ('h', 'Deadlines'),
        ('p', 'The register calculates: the summary resolution deadline (close of the third business day after receipt); the eight-week final response deadline; and, for payment services complaints, the 15-business-day deadline (35 business days in exceptional circumstances). The business-day formulas do not exclude bank holidays – check dates close to bank holidays. Delete the payment services columns if you do not provide payment services.'),
        ('h', 'Root causes and reporting'),
        ('p', 'Record a root cause for every complaint. The Summary sheet counts complaints by category and outcome for your management information, Consumer Duty monitoring and the FCA complaints return (from 1 January 2027, the consolidated complaints return under PS25/19).'),
        ('h', 'Data protection complaints'),
        ('p', 'Since 19 June 2026 data protection complaints must be acknowledged within 30 days. Tick the data protection column so the acknowledgement deadline is calculated.'),
    ])
    b.add_list('chan', ['Telephone', 'Email', 'Letter', 'Web form', 'In person', 'Social media', 'Via third party / CMC'])
    b.add_list('act', ACT_NAMES)
    b.add_list('cat', ['Advice / suitability', 'Information / disclosure', 'Affordability / lending decision', 'Fees and charges', 'Arrears / collections', 'Customer service / delays', 'Staff conduct', 'Fraud / scams', 'Data protection', 'Product performance', 'Claims (insurance)', 'Other'])
    b.add_list('out', ['Resolved by summary resolution communication', 'Upheld', 'Partly upheld', 'Not upheld', 'Forwarded to another firm (DISP 1.7)', 'Withdrawn'])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('fos', ['Not referred', 'Referred – awaiting outcome', 'Ombudsman upheld', 'Ombudsman not upheld', 'Settled'])
    cols = ['Ref.', 'Date received', 'Received by', 'Channel', 'Complainant name', 'Customer / representative', 'Regulated activity / product', 'Complaint category', 'Summary of complaint', 'Vulnerability / adjustments', 'Data protection complaint?', 'Payment services (PSRs Parts 6/7) complaint?',
            'Acknowledged on', 'Summary resolution deadline (3 business days)', 'Final response deadline (8 weeks)', 'PSR deadline (15 business days)', 'PSR exceptional deadline (35 business days)', 'Data protection acknowledgement deadline (30 days)',
            'Handler', 'Date of final response / 8-week letter', 'Outcome', 'Redress (£ and type)', 'Date redress paid', 'Forwarded to (firm)', 'Root cause', 'Wider action taken (Corrective Action Log ref.)', 'Ombudsman referral', 'Ombudsman outcome and date']
    widths = [9, 11, 14, 12, 18, 14, 18, 18, 36, 20, 11, 12, 11, 13, 12, 12, 12, 13, 14, 13, 20, 14, 11, 16, 26, 22, 18, 20]
    ws = b.sheet('Complaints Register', widths, 'Complaints Register')
    r = b.note(ws, 3, 'Record every complaint by the end of the next business day. Deadlines calculate automatically from the date received. Delete columns L, P and Q if you do not provide payment services.', len(cols))
    r = b.header(ws, r, cols, height=60, groups=[('Complaint details', 12), ('Deadlines (calculated)', 6), ('Response and outcome', 6), ('Learning and Ombudsman', 4)])
    fx = {14: '=IF(B{r}="","",WORKDAY(B{r},3))', 15: '=IF(B{r}="","",B{r}+56)', 16: '=IF(OR(B{r}="",L{r}<>"Yes"),"",WORKDAY(B{r},15))', 17: '=IF(OR(B{r}="",L{r}<>"Yes"),"",WORKDAY(B{r},35))', 18: '=IF(OR(B{r}="",K{r}<>"Yes"),"",B{r}+30)'}
    end = b.rows(ws, r, 150, len(cols), height=24, formulas=fx, first_ref='CR-')
    for key, col in [('chan', 4), ('act', 7), ('cat', 8), ('yn', 11), ('yn', 12), ('out', 21), ('fos', 27)]:
        b.validate(ws, key, col, r, end - 1)
    for col in (2, 13, 14, 15, 16, 17, 18, 20, 23):
        for rr in range(r, end):
            ws.cell(rr, col).number_format = 'DD/MM/YYYY'
    ws.conditional_formatting.add(f'O{r}:O{end - 1}', FormulaRule(formula=[f'AND(O{r}<>"",T{r}="",O{r}<TODAY())'], fill=PatternFill('solid', fgColor='F4B6B6')))
    # Summary
    ws2 = b.sheet('Summary', [34, 14, 14, 14, 14], 'Complaints Summary (management information)', landscape=False)
    r = b.header(ws2, 4, ['Complaint category', 'Received', 'Upheld / partly upheld', 'Not upheld', 'Referred to Ombudsman'], height=36)
    cats = ['Advice / suitability', 'Information / disclosure', 'Affordability / lending decision', 'Fees and charges', 'Arrears / collections', 'Customer service / delays', 'Staff conduct', 'Fraud / scams', 'Data protection', 'Product performance', 'Claims (insurance)', 'Other']
    rng = "'Complaints Register'!"
    for cat in cats:
        vals = [cat, f'=COUNTIF({rng}H:H,A{r})', f'=COUNTIFS({rng}H:H,A{r},{rng}U:U,"Upheld")+COUNTIFS({rng}H:H,A{r},{rng}U:U,"Partly upheld")', f'=COUNTIFS({rng}H:H,A{r},{rng}U:U,"Not upheld")', f'=COUNTIFS({rng}H:H,A{r},{rng}AA:AA,"<>Not referred",{rng}AA:AA,"<>")']
        r = b.data_row(ws2, r, vals, 5)
    b.data_row(ws2, r, ['Total', f'=SUM(B5:B{r - 1})', f'=SUM(C5:C{r - 1})', f'=SUM(D5:D{r - 1})', f'=SUM(E5:E{r - 1})'], 5, bold=True)
    b.save(f'{OUT}/Complaints_Register_v2.0_Template.xlsx')


# =====================================================================
# 6. Financial Promotions Log (new) – log + approval checklist
# =====================================================================
FP_ACT_CHECKS = {
    'Credit broking': ['Broker status shown: we are a credit broker, not a lender; extent of lender panel; any fee disclosed (CONC 3).', 'Representative example included, and more prominent than the trigger, where a rate or cost figure is shown; representative APR meets the 51% test and evidence is on file (CONC 3.5).', 'No “pre-approved”, “guaranteed” or “regardless of credit history” claims; HCSTC risk warning included if applicable.'],
    'Debt adjusting and debt counselling': ['States that free debt advice is available and how to find it (e.g. MoneyHelper).', 'No misleading claims that debts will be written off, that the solution is government-backed or free.', 'Fees disclosed clearly; lead source authorised and checked.'],
    'Debt collecting and debt administration': ['Collection communication does not misrepresent authority or legal position, or threaten action we cannot or will not take (CONC 7).', 'Signposts free debt advice.'],
    'Credit information services and credit references': ['Does not overstate what the service can do (e.g. “fix your credit score”).', 'Statutory free access made clear; free trial / subscription price, start date and cancellation shown with equal prominence.'],
    'Peer-to-peer lending platforms': ['Prescribed restricted mass market investment risk warning included and prominent, with link to risk summary (COBS 4.12A).', 'No incentives to invest; audience restricted to categorised investors; direct offer promotions only after appropriateness, with 24-hour cooling-off and personalised risk warning.', 'Borrower promotions: CONC 3 representative example where a rate or cost is shown.'],
    'Consumer credit lending': ['Representative example included where a rate or cost figure is shown; representative APR meets the 51% test and evidence is on file (CONC 3.5).', 'No “pre-approved” or “regardless of credit history” claims; HCSTC risk warning if applicable; deferred payment credit not presented as risk-free.', 'Promotions communicated by brokers / merchants: approved content only; s21 approver regime position checked.'],
    'Exercising a lender’s rights and duties': ['Settlement or arrangement offers clear on terms and credit file effect; no false urgency.', 'Communication clear about who we are; does not misrepresent legal position (CONC 7).'],
    'Consumer hire': ['States the customer will not own the goods; total cost, hire period and end-of-hire / damage / excess charges shown fairly (CONC 3).'],
    'Insurance distribution': ['Cover, key exclusions and limitations presented fairly (ICOBS 2.2).', '“From £x” prices achievable by a meaningful proportion of customers; comparisons fair.', 'Home / motor: consistent with pricing practices rules (ICOBS 6B); add-ons not presented as compulsory.'],
    'Mortgage intermediation': ['Repossession risk warning included (and debt consolidation warning if applicable) (MCOB 3A).', 'Representative example with APRC where a rate or cost figure is shown (MCOB 3A).', 'Fee and whole-of-market / limited panel status stated.'],
    'Payment services and e-money': ['Does not call e-money or safeguarded funds a deposit or bank account, or imply FSCS protection.', 'Fees and exchange rates (including mark-ups) shown clearly; speed and security not overstated.'],
    'Investment advice and arranging': ['Fair, clear and not misleading with prominent risks; past / future performance and comparison rules met (COBS 4).', 'High-risk investment restrictions and prescribed risk warnings met (COBS 4.12A / 4.12B) if applicable.', 'Nature of service (independent / restricted) and charges stated.'],
}


def finprom_log():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Financial Promotions Log – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', '1. Financial Promotions Log – one line for every financial promotion (and version) approved, including social media posts, influencer and affiliate content, and scripts for real-time promotions. This is the “Financial Promotions Log” referred to in the Financial Promotions Policy.'),
        ('p', '2. Approval Checklist – complete a copy for each promotion before approval. This is the “Financial Promotion Checklist” referred to in the policy. Keep the completed checklist with the promotion.'),
        ('h', 'Records'),
        ('p', 'Keep a copy of each promotion exactly as it appeared (screenshots for social media), its approval, the evidence for its claims (for example the basis of a representative APR), and when it was withdrawn, for at least [Insert period].'),
        ('watch', 'The DMCC Act 2024 subscription contracts regime is now expected to commence in spring 2027 and will affect promotions of subscription services.'),
    ])
    b.add_list('chan', ['Website / app', 'Social media – organic', 'Social media – paid', 'Influencer / affiliate', 'Third-party platform', 'Email / SMS', 'Print / point of sale', 'Telephone / face-to-face script', 'Video', 'Other'])
    b.add_list('act', ACT_NAMES)
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('status', ['Draft', 'Approved – live', 'Approved – not yet live', 'Withdrawn', 'Rejected'])
    cols = ['Ref.', 'Promotion title / description', 'Version', 'Regulated activity / product', 'Channel(s)', 'Target audience', 'Created by', 'Approver', 'Date approved', 'Approval checklist completed?', 'Key rules applied (e.g. CONC 3.5, MCOB 3A, COBS 4.12A)', 'Evidence for claims (location)', 'Communicated by a third party? (name)', 'Date first used', 'Review due', 'Date withdrawn', 'Reason withdrawn', 'Status', 'Copy / screenshot location']
    widths = [9, 30, 8, 20, 18, 20, 14, 14, 11, 12, 24, 20, 18, 11, 11, 11, 18, 14, 22]
    ws = b.sheet('Financial Promotions Log', widths, 'Financial Promotions Log')
    r = b.note(ws, 3, 'Nothing is published until it has been approved and logged. Review live promotions at least [Insert frequency] – the review date turns red when overdue.', len(cols))
    r = b.header(ws, r, cols, height=55)
    end = b.rows(ws, r, 150, len(cols), height=22, first_ref='FP-')
    for key, col in [('act', 4), ('chan', 5), ('yn', 10), ('status', 18)]:
        b.validate(ws, key, col, r, end - 1)
    for col in (9, 14, 15, 16):
        for rr in range(r, end):
            ws.cell(rr, col).number_format = 'DD/MM/YYYY'
    ws.conditional_formatting.add(f'O{r}:O{end - 1}', FormulaRule(formula=[f'AND(O{r}<>"",P{r}="",O{r}<TODAY())'], fill=PatternFill('solid', fgColor='F4B6B6')))
    # Checklist
    cols = ['Area', 'Check', 'Met? (Yes / No / N/A)', 'Comments / changes required']
    widths = [22, 80, 13, 40]
    ws = b.sheet('Approval Checklist', widths, 'Financial Promotion Approval Checklist', 'Promotion ref.: [Insert]   |   Approver: [Insert name]   |   Date: [Insert date]')
    r = b.note(ws, 3, 'Complete before approving any promotion. Every applicable check must be “Yes” before approval.', len(cols))
    r = b.header(ws, r, cols, height=30)
    first = r
    universal = [
        ('Is it a promotion?', 'The communication is (or contains) an invitation or inducement to engage in investment activity, or is a customer communication that must still be clear, fair and not misleading.'),
        ('Section 21', 'We are permitted to communicate / approve it; if approving for an unauthorised person, we hold the s21 approver permission or an exemption applies.'),
        ('Target audience', 'Target audience identified and consistent with the product’s target market; not aimed at customers likely to be harmed.'),
        ('Fair, clear, not misleading', 'Accurate, evidenced and balanced; benefits not emphasised without equally prominent risks, costs and limitations.'),
        ('Prominence', 'Required information and warnings are prominent and in the promotion itself (not only behind a link or in a comment).'),
        ('Plain language', 'Plain English; technical terms explained; suitable for customers with characteristics of vulnerability.'),
        ('No pressure', 'No false urgency, countdown timers or exploitation of behavioural biases.'),
        ('Firm identified', 'Firm name and, where required, FCA status shown.'),
        ('Channel rules', 'Meets the channel standards in the policy (e.g. FG24/1 for social media: each post compliant on its own; adverts labelled).'),
        ('Consumer Duty', 'Supports consumer understanding (PRIN 2A.5) and is consistent with the product’s fair value assessment.'),
        ('Records', 'Final version, evidence and approval saved; entered in the Financial Promotions Log with review date.'),
    ]
    for area, chk in universal:
        r = b.data_row(ws, r, [area, chk], len(cols), height=36)
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        for chk in FP_ACT_CHECKS[act]:
            r = b.data_row(ws, r, ['Activity rules', chk], len(cols), height=36)
    r = b.data_row(ws, r, ['Decision', 'Approved / rejected: [Insert]   Approver signature: [Insert]'], len(cols), height=30, bold=True)
    b.validate(ws, 'ynna', 3, first, r - 1)
    b.save(f'{OUT}/Financial_Promotions_Log_v2.0_Template.xlsx')


# =====================================================================
# 7. Data Protection Registers (new)
# =====================================================================
def data_registers():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Data Protection Registers – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', 'The registers referred to in the Data Processing Policy: Record of Processing Activities; Data Rights Request Log; Data Breach Log; Data Protection Complaints Log; Data Retention Schedule; and DPIA Register.'),
        ('h', 'Key deadlines'),
        ('p', 'Rights requests: respond within one month (extendable by two months for complex requests; paused while waiting for information needed to identify the person or clarify the request). Personal data breaches: notify the ICO within 72 hours where there is a risk to individuals. Data protection complaints: acknowledge within 30 days (from 19 June 2026). The logs calculate these dates.'),
        ('watch', 'The Data (Use and Access) Act 2025 changes are being commenced in stages and ICO guidance is still being updated. Check the ICO website for the current position.'),
    ])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('right', ['Access (SAR)', 'Rectification', 'Erasure', 'Restriction', 'Portability', 'Objection', 'Objection – direct marketing', 'Automated decision-making'])
    b.add_list('basis', ['Contract', 'Legal obligation', 'Legitimate interests', 'Recognised legitimate interests', 'Consent', 'Vital interests', 'Public task'])
    b.add_list('risk', ['No risk – record only', 'Risk – notify ICO', 'High risk – notify ICO and individuals'])
    b.add_list('act', ACT_NAMES)
    # ROPA
    cols = ['Ref.', 'Processing activity', 'Regulated activity', 'Purpose', 'Data subjects', 'Personal data categories', 'Special category / criminal offence data? (condition relied on)', 'Lawful basis', 'Source of data', 'Recipients / sharing', 'International transfers (safeguard)', 'Processors used', 'Retention period', 'Security measures', 'DPIA needed? (ref.)', 'Owner', 'Last reviewed']
    widths = [7, 24, 18, 24, 16, 22, 22, 16, 16, 20, 18, 16, 14, 18, 12, 13, 11]
    ws = b.sheet('ROPA', widths, 'Record of Processing Activities')
    r = b.note(ws, 3, 'One line per processing activity (Article 30 UK GDPR). Starter lines are examples – adapt them to your own processing and add activity-specific processing from section 8 of the Data Processing Policy.', len(cols))
    r = b.header(ws, r, cols, height=55)
    first = r
    starters = [
        ['P01', 'Customer onboarding and due diligence', 'All activities', 'Verify identity; prevent fraud and financial crime', 'Customers, guarantors, representatives', 'Identity, contact, verification documents', '[Insert or “None”]', 'Legal obligation', 'Customer; verification provider', '[Insert]', '[Insert]', '[Insert]', '5 years after relationship ends', '[Insert]', '[Insert]', '[Insert]'],
        ['P02', 'Providing our regulated service', 'All activities', 'Deliver the service the customer asked for', 'Customers', '[Insert]', '[Insert]', 'Contract', 'Customer', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]'],
        ['P03', 'Complaints handling', 'All activities', 'Investigate and resolve complaints', 'Complainants, representatives', 'Contact, complaint details', '[Insert]', 'Legal obligation', 'Complainant', 'Financial Ombudsman Service; other firms (DISP 1.7)', 'None', '[Insert]', '[Insert – at least 3 years]', '[Insert]', 'No', '[Insert]'],
        ['P04', 'Customer vulnerability records', 'All activities', 'Support customers’ needs', 'Customers', 'Support needs; health (where disclosed)', 'Yes – explicit consent / DPA 2018 Sch. 1 condition', 'Legitimate interests / legal obligation', 'Customer', 'Other firms with consent', 'None', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]'],
        ['P05', 'Marketing', 'All activities', 'Promote our services', 'Customers, prospects', 'Contact, preferences', 'No', 'Consent / legitimate interests (soft opt-in)', 'Customer', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]'],
        ['P06', 'Employment and HR', 'All activities', 'Employ and manage staff', 'Staff, applicants', 'HR, payroll, training, F&P records', 'Yes – [Insert condition]', 'Contract / legal obligation', 'Staff; referees; DBS', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]', '[Insert]'],
    ]
    for s in starters:
        r = b.data_row(ws, r, s, len(cols), height=40)
    end = b.rows(ws, r, 20, len(cols), height=24)
    b.validate(ws, 'act', 3, first, end - 1)
    # Rights requests
    cols = ['Ref.', 'Date received', 'Requester name', 'Customer / staff / other', 'Right exercised', 'Channel received', 'Identity verified on', 'Clarification requested? (date – clock paused)', 'Date clarification received', 'Response deadline (1 month)', 'Extended? (reason, new deadline)', 'Handler', 'Date responded', 'Outcome (complied / partly / refused – reason)', 'Exemptions applied', 'Complaint to ICO?']
    widths = [9, 11, 18, 14, 18, 13, 12, 16, 12, 13, 18, 13, 12, 24, 18, 11]
    ws = b.sheet('Data Rights Request Log', widths, 'Data Rights Request Log')
    r = b.note(ws, 3, 'Requests can be made in any form, including verbally or by social media. The one-month deadline is calculated from the date received (or from the date clarification was received, if later). Check the result against the ICO’s calendar-month rule.', len(cols))
    r = b.header(ws, r, cols, height=55)
    end = b.rows(ws, r, 60, len(cols), height=22, first_ref='DR-', formulas={10: '=IF(B{r}="","",EDATE(MAX(B{r},N(I{r})),1))'})
    b.validate(ws, 'right', 5, r, end - 1); b.validate(ws, 'yn', 16, r, end - 1)
    for col in (2, 7, 9, 10, 13):
        for rr in range(r, end):
            ws.cell(rr, col).number_format = 'DD/MM/YYYY'
    # Breach log
    cols = ['Ref.', 'Date and time of breach', 'Date and time we became aware', 'Reported by', 'Description of breach', 'Personal data involved', 'Number of individuals affected', 'Containment actions', 'Risk assessment', 'ICO notification deadline (72 hours from awareness)', 'ICO notified? (date / time and ref.)', 'Reason not notified', 'Individuals informed? (date)', 'FCA notification considered? (Compliance Breach Log ref.)', 'Root cause', 'Remedial and preventive action', 'Owner', 'Date closed']
    widths = [9, 14, 14, 13, 30, 20, 11, 24, 20, 15, 16, 18, 13, 18, 22, 26, 13, 11]
    ws = b.sheet('Data Breach Log', widths, 'Data Breach Log')
    r = b.note(ws, 3, 'Record every personal data breach, including those you do not report. Enter the date and time you became aware as a date-time (e.g. 25/09/2026 14:30) – the 72-hour ICO deadline calculates automatically.', len(cols))
    r = b.header(ws, r, cols, height=55)
    end = b.rows(ws, r, 60, len(cols), height=24, first_ref='DB-', formulas={10: '=IF(C{r}="","",C{r}+3)'})
    b.validate(ws, 'risk', 9, r, end - 1)
    for col in (2, 3, 10):
        for rr in range(r, end):
            ws.cell(rr, col).number_format = 'DD/MM/YYYY HH:MM'
    # DP complaints
    cols = ['Ref.', 'Date received', 'Complainant', 'Channel', 'Summary', 'Acknowledgement deadline (30 days)', 'Date acknowledged', 'Enquiries made', 'Handler', 'Date outcome sent', 'Outcome', 'Also a DISP complaint? (Complaints Register ref.)', 'Told of right to complain to ICO?']
    widths = [9, 11, 18, 12, 32, 14, 12, 26, 13, 12, 26, 18, 13]
    ws = b.sheet('DP Complaints Log', widths, 'Data Protection Complaints Log')
    r = b.note(ws, 3, 'Since 19 June 2026 we must have a process for data protection complaints: acknowledge within 30 days, respond without undue delay, and tell the complainant the outcome. If the complaint is also about our financial services, record it in the Complaints Register too and meet the shorter deadline.', len(cols))
    r = b.header(ws, r, cols, height=55)
    end = b.rows(ws, r, 40, len(cols), height=22, first_ref='DC-', formulas={6: '=IF(B{r}="","",B{r}+30)'})
    b.validate(ws, 'yn', 13, r, end - 1)
    # Retention schedule
    cols = ['Record type', 'Regulated activity', 'Retention period', 'Trigger (start of period)', 'Reason / legal basis', 'Deletion method', 'Owner']
    widths = [34, 20, 22, 24, 36, 20, 14]
    ws = b.sheet('Data Retention Schedule', widths, 'Data Retention Schedule')
    r = b.note(ws, 3, 'Keep this consistent with section 14 and section 8 of the Data Processing Policy. Replace every placeholder with your own period and delete rows for activities you do not carry on.', len(cols))
    r = b.header(ws, r, cols, height=36)
    sched = [
        ['Customer due diligence records', 'All activities', '5 years', 'End of relationship / occasional transaction', 'Money Laundering Regulations (where they apply); delete afterwards unless law, proceedings or consent allow longer', 'Secure deletion', '[Insert]'],
        ['Complaint records', 'All activities', '[Insert – at least 3 years]', 'Date complaint received', 'DISP 1.9; Ombudsman time limits', 'Secure deletion', '[Insert]'],
        ['General customer and contract records', 'All activities', '[Insert, e.g. 6 years]', 'End of relationship', 'Limitation Act 1980; SYSC 9', 'Secure deletion', '[Insert]'],
        ['Financial promotion records', 'All activities', '[Insert]', 'Promotion last used', 'FCA rules for our activity', 'Secure deletion', '[Insert]'],
        ['Marketing consent records', 'All activities', 'While consent relied on + [Insert]', 'Consent withdrawn / last used', 'UK GDPR accountability; PECR', 'Secure deletion', '[Insert]'],
        ['Employee records', 'All activities', '[Insert, e.g. 6 years]', 'Employment ends', 'Employment law; limitation', 'Secure deletion', '[Insert]'],
        ['Training and competence records', 'All activities', '[Insert] (indefinitely for pension transfer specialists)', 'Person leaves', 'TC; SYSC', 'Secure deletion', '[Insert]'],
        ['CCTV', 'All activities', '[Insert, e.g. 30 days]', 'Recording', 'Security; proportionality', 'Automatic overwrite', '[Insert]'],
        ['Investment advice suitability records', 'Investment advice and arranging', 'As COBS requires (indefinitely for pension transfers)', 'Advice given', 'COBS 9 / 9A', 'Secure deletion', '[Insert]'],
        ['Recorded communications (MiFID)', 'Investment advice and arranging', '5 years', 'Date of recording', 'COBS 11.8', 'Secure deletion', '[Insert]'],
        ['Mortgage advice records', 'Mortgage intermediation', '[Insert – at least the MCOB minimum]', 'Advice given', 'MCOB 4; limitation', 'Secure deletion', '[Insert]'],
        ['Motor finance commission records', 'Credit broking / Consumer credit lending', 'Retain while FCA retention requirements apply', 'N/A', 'FCA motor finance requirements', 'Do not delete', '[Insert]'],
    ]
    for s in sched:
        r = b.data_row(ws, r, s, len(cols), height=32)
    b.rows(ws, r, 10, len(cols), height=22)
    # DPIA
    cols = ['DPIA ref.', 'Processing assessed', 'Why high risk', 'Date completed', 'Key risks identified', 'Measures to reduce risk', 'Residual risk acceptable?', 'ICO consulted? (date)', 'Approved by', 'Next review']
    widths = [9, 26, 22, 12, 28, 28, 12, 14, 14, 12]
    ws = b.sheet('DPIA Register', widths, 'DPIA Register')
    r = b.note(ws, 3, 'List every data protection impact assessment. A DPIA is required before high-risk processing (for example credit scoring, large-scale special category data, systematic monitoring, new technology including AI tools).', len(cols))
    r = b.header(ws, r, cols, height=40)
    end = b.rows(ws, r, 20, len(cols), height=22, first_ref='DPIA-')
    b.validate(ws, 'yn', 7, r, end - 1)
    b.save(f'{OUT}/Data_Protection_Registers_v2.0_Template.xlsx')


# =====================================================================
# 8. Conflicts and Gifts Registers (new)
# =====================================================================
def conflicts_gifts():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Conflicts and Gifts Registers – how to use', STANDARD_GUIDE + [
        ('h', 'Conflicts of Interest Register'),
        ('p', 'Record every actual or potential conflict between the interests of the firm (including staff, owners and partners) and its customers, or between customers. Typical conflicts include commission and remuneration arrangements, volume incentives, close relationships with lenders, insurers or introducers, and staff outside interests. Record how each is managed and review the register at least [Insert frequency]. Disclosure to customers is a last resort, not a substitute for managing the conflict.'),
        ('h', 'Gifts and Hospitality Register'),
        ('p', 'Record all gifts and hospitality given or received above [Insert value, e.g. £50] and whether they were approved in advance, as the Anti-Money Laundering and Financial Crime Policy (section 9.2) requires. Record declined offers too.'),
    ])
    b.add_list('ctype', ['Commission / remuneration from providers', 'Staff pay and incentives', 'Relationship with lender / insurer / introducer', 'Personal or outside interest', 'Ownership / group', 'Between customers', 'Other'])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('gr', ['Given', 'Received', 'Offered and declined'])
    b.add_list('status', ['Open – managed', 'Closed', 'Under review'])
    cols = ['Ref.', 'Date identified', 'Description of conflict', 'Type', 'Parties involved', 'Risk of harm to customers', 'How it is managed / mitigated', 'Disclosed to customers? (how)', 'Owner', 'Last reviewed', 'Next review', 'Status']
    widths = [8, 11, 34, 22, 20, 24, 34, 18, 14, 11, 11, 14]
    ws = b.sheet('Conflicts of Interest Register', widths, 'Conflicts of Interest Register')
    r = b.note(ws, 3, 'Record actual and potential conflicts. Starter lines are examples – adapt or delete.', len(cols))
    r = b.header(ws, r, cols, height=40)
    first = r
    r = b.data_row(ws, r, ['COI-001', None, 'We receive commission from [Insert providers] which may vary between providers and could influence which product we recommend or introduce.', 'Commission / remuneration from providers', '[Insert]', 'Customers introduced to a more expensive or less suitable product.', 'Commission does not vary by product in a way that incentivises poor outcomes; commission disclosed as rules require; file reviews.', '[Insert]', '[Insert]'], len(cols), height=60)
    r = b.data_row(ws, r, ['COI-002', None, 'Staff pay, bonuses or targets are linked to sales volumes.', 'Staff pay and incentives', 'Staff', 'Pressure selling; poor outcomes.', 'Incentives balanced with quality measures; file reviews; no incentive overrides the right to pause or stop a sale.', 'No', '[Insert]'], len(cols), height=60)
    end = b.rows(ws, r, 30, len(cols), height=24, first_ref=None)
    for key, col in [('ctype', 4), ('status', 12)]:
        b.validate(ws, key, col, first, end - 1)
    cols = ['Ref.', 'Date', 'Given / received', 'Staff member', 'Counterparty (organisation and person)', 'Description', 'Estimated value (£)', 'Approved in advance? (by whom)', 'Reason / business purpose', 'Any concern about influence?']
    widths = [8, 11, 14, 16, 26, 30, 11, 20, 26, 16]
    ws = b.sheet('Gifts and Hospitality Register', widths, 'Gifts and Hospitality Register')
    r = b.note(ws, 3, 'Record all gifts and hospitality above [Insert value] given, received or declined. Cash or cash-equivalent gifts must always be refused.', len(cols))
    r = b.header(ws, r, cols, height=40)
    end = b.rows(ws, r, 60, len(cols), height=22, first_ref='GH-')
    b.validate(ws, 'gr', 3, r, end - 1)
    b.save(f'{OUT}/Conflicts_and_Gifts_Registers_v2.0_Template.xlsx')


# =====================================================================
# 9. AML Business-Wide Risk Assessment (new)
# =====================================================================
BWRA_FACTORS = [
    ('Customers', 'Customer types (consumers, sole traders, companies, trusts)', '[Describe your customers]'),
    ('Customers', 'Politically exposed persons', '[Expected number / likelihood]'),
    ('Customers', 'Customers acting for others / third-party payers', '[Describe]'),
    ('Customers', 'Customers with characteristics of vulnerability (risk of exploitation / money mules)', '[Describe]'),
    ('Products and services', 'Nature of our products and services (value, speed, anonymity, ability to move value)', '[Describe]'),
    ('Products and services', 'Transaction sizes and payment methods (cash, third-party, overpayments, early settlement)', '[Describe]'),
    ('Delivery channels', 'Face-to-face, remote/online, introducers, brokers or appointed representatives', '[Describe]'),
    ('Delivery channels', 'Reliance on third parties for CDD', '[Describe]'),
    ('Geography', 'Where customers live and where funds come from or go to', '[Describe]'),
    ('Geography', 'Links to FATF “call for action” or increased-monitoring jurisdictions', '[Describe]'),
    ('Sanctions', 'Exposure to designated persons or sanctioned jurisdictions', '[Describe]'),
    ('Fraud', 'Application, identity, payment and internal fraud', '[Describe]'),
    ('Other', 'Bribery and facilitation of tax evasion', '[Describe]'),
]
BWRA_ACT = {
    'Credit broking': 'Application fraud; third-party funded deposits; introducer misconduct.',
    'Debt adjusting and debt counselling': 'Client money misuse; third-party payments into plans; coercion.',
    'Debt collecting and debt administration': 'Lump-sum / third-party settlements; overpayment refunds; impersonation.',
    'Credit information services and credit references': 'Identity theft; data manipulation; insider misuse.',
    'Peer-to-peer lending platforms': 'Criminal funds placed by investors; borrower fraud; collusion; pooled accounts.',
    'Consumer credit lending': 'Loan-back and early settlement; synthetic identities; money mules; merchant fraud.',
    'Exercising a lender’s rights and duties': 'Inherited CDD gaps; unexplained settlements; portfolio sellers.',
    'Consumer hire': 'Fraudulent hire of high-value goods; onward sale or export; refunds.',
    'Insurance distribution': 'Ghost broking; premium and claims fraud; single premiums and early surrender (life).',
    'Mortgage intermediation': 'Mortgage fraud; criminal deposits; property used to launder funds.',
    'Payment services and e-money': 'Rapid movement of funds; mules; APP fraud; account takeover; high-risk corridors.',
    'Investment advice and arranging': 'Criminal lump sums; unexplained wealth; early encashment; complex structures; PEPs.',
}


def bwra():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Business-Wide Risk Assessment – how to use', STANDARD_GUIDE + [
        ('h', 'What this is'),
        ('p', 'The written Business-Wide Risk Assessment (BWRA) referred to in section 5 of the Anti-Money Laundering and Financial Crime Policy. Firms within the Money Laundering Regulations must have one (regulation 18). Firms outside the MLRs use it as the foundation of the controls SYSC 6.1.1R requires.'),
        ('h', 'How to complete it'),
        ('p', 'For each risk factor, describe your exposure, score inherent risk (Low / Medium / High), list your controls, and score residual risk. Take account of the UK National Risk Assessment of money laundering and terrorist financing, FCA publications and relevant JMLSG sector guidance. Complete the Conclusion sheet and have it approved by senior management. Review at least annually and whenever your business changes materially.'),
    ])
    b.add_list('lmh', ['Low', 'Medium', 'High'])
    cols = ['Area', 'Risk factor', 'Our exposure (describe)', 'Inherent risk', 'Controls in place', 'Residual risk', 'Rationale', 'Actions required']
    widths = [18, 38, 38, 11, 38, 11, 30, 28]
    ws = b.sheet('Risk Factors', widths, 'Business-Wide Risk Assessment – Risk Factors')
    r = b.note(ws, 3, 'Complete every line. Activity-specific typologies sit under green OPTION rows – keep those for your activities.', len(cols))
    r = b.header(ws, r, cols, height=36)
    first = r
    for area, factor, exp in BWRA_FACTORS:
        r = b.data_row(ws, r, [area, factor, exp], len(cols), height=40)
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        r = b.data_row(ws, r, ['Activity typologies', f'{act}: {BWRA_ACT[act]}', '[Describe]'], len(cols), height=40)
    last = r - 1
    for col in (4, 6):
        b.validate(ws, 'lmh', col, first, last)
        rng = f'{L(col)}{first}:{L(col)}{last}'
        for val, colr in (('High', 'F4B6B6'), ('Medium', 'FFE699'), ('Low', 'C6E0B4')):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L(col)}{first}="{val}"'], fill=PatternFill('solid', fgColor=colr)))
    ws2 = b.sheet('Conclusion', [40, 80], 'Business-Wide Risk Assessment – Conclusion', landscape=False)
    r = 4
    for label, val in [('Overall inherent risk', '[Low / Medium / High]'), ('Overall residual risk', '[Low / Medium / High]'), ('Key risks', '[Insert]'), ('How this assessment affects our EDD triggers', '[Insert]'), ('How it affects our monitoring and training', '[Insert]'), ('Sources considered', 'UK National Risk Assessment; FCA publications; JMLSG sector guidance; [Insert others]'), ('Prepared by (MLRO)', '[Insert name and date]'), ('Approved by senior management', '[Insert name and date]'), ('Next review date', '[Insert date]')]:
        r = b.data_row(ws2, r, [label, val], 2, height=32, bold=False)
        ws2.cell(r - 1, 1).font = b.f(b=True, color=GREEN_TEXT)
    b.save(f'{OUT}/AML_Business_Wide_Risk_Assessment_v2.0_Template.xlsx')


# =====================================================================
# 10. Compliance Monitoring Tracker (new)
# =====================================================================
def monitoring_tracker():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Compliance Monitoring Tracker – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', '1. Monitoring Schedule – tracks each check in your Compliance Monitoring Plan (Part A and your Part B options): when it is due, when it was done, and the result.'),
        ('p', '2. File Review Record – records the result of each file reviewed, so findings can be evidenced and trends identified. Use it with the checklist for the relevant area.'),
        ('p', '3. Corrective Action Log – every Red or Amber finding, breach or significant complaint theme, tracked to verified closure. This is the “Corrective Action Log” referred to in the Compliance Monitoring Plan.'),
        ('p', '4. Compliance Universe Register – the requirements that apply to your firm, and where each is covered. This is the “Compliance Universe Register” referred to in the Compliance Monitoring Programme.'),
        ('p', 'Breaches are recorded in the Compliance Breach Log (a separate workbook) and cross-referenced here.'),
    ])
    b.add_list('rag', ['Green – satisfactory', 'Amber – weakness', 'Red – breach / customer harm'])
    b.add_list('freq', ['Monthly', 'Quarterly', 'Six-monthly', 'Annually', 'Ad hoc'])
    b.add_list('src', ['Compliance monitoring', 'Breach', 'Complaint theme', 'Incident', 'External review / audit', 'FCA feedback', 'Other'])
    b.add_list('status', ['Open', 'In progress', 'Awaiting verification', 'Closed – verified', 'Overdue'])
    b.add_list('yn', YN); b.add_list('ynna', YNNA)
    b.add_list('applies', ['Yes', 'No', 'Partly'])
    # Schedule
    cols = ['Plan ref. (e.g. A1, BA2)', 'Area', 'Risk rating', 'Frequency', 'Owner', 'Period covered', 'Due date', 'Date completed', 'Files / items reviewed', 'Result', 'Findings (summary)', 'Corrective Action Log refs.', 'Reported to board / owner (date)']
    widths = [12, 26, 10, 12, 14, 14, 11, 11, 11, 20, 36, 16, 14]
    ws = b.sheet('Monitoring Schedule', widths, 'Compliance Monitoring Schedule Tracker')
    r = b.note(ws, 3, 'Enter one line per check per period, using the references from your Compliance Monitoring Plan. Overdue checks turn red.', len(cols))
    r = b.header(ws, r, cols, height=40)
    first = r
    for ref, area, risk, freq in [('A1', 'Consumer Duty – customer outcomes', 'High', 'Quarterly'), ('A2', 'Financial promotions and communications', 'High', 'Monthly'), ('A3', 'Complaints handling', 'Medium', 'Quarterly'), ('A4', 'Customers with characteristics of vulnerability', 'High', 'Quarterly'), ('A5', 'Financial crime', 'Medium', 'Quarterly'), ('A6', 'Training and competence', 'Medium', 'Six-monthly'), ('A7', 'Governance and SM&CR', 'Low', 'Annually'), ('A8', 'Data protection and record keeping', 'Medium', 'Quarterly'), ('A9', 'Conflicts of interest and remuneration', 'Medium', 'Six-monthly'), ('A10', 'Outsourcing and third parties', 'Low', 'Annually'), ('A11', 'Regulatory reporting and notifications', 'Medium', 'Quarterly'), ('A12', 'Financial resources and wind-down', 'Medium', 'Quarterly'), ('A13', 'Appointed representatives [delete if none]', 'High', 'Quarterly')]:
        r = b.data_row(ws, r, [ref, area, risk, freq, '[Insert owner]'], len(cols), height=22)
    r = b.data_row(ws, r, ['[Insert Part B refs]', '[Insert activity-specific areas from Part B of your plan]'], len(cols), height=22)
    end = b.rows(ws, r, 60, len(cols), height=20)
    b.validate(ws, 'freq', 4, first, end - 1); b.validate(ws, 'rag', 10, first, end - 1)
    b.rag(ws, f'J{first}:J{end - 1}')
    ws.conditional_formatting.add(f'G{first}:G{end - 1}', FormulaRule(formula=[f'AND(G{first}<>"",H{first}="",G{first}<TODAY())'], fill=PatternFill('solid', fgColor='F4B6B6')))
    # File review record
    cols = ['Review ref.', 'Plan ref.', 'Date reviewed', 'Reviewer', 'Customer / file ref.', 'Staff member responsible', 'Why selected (random / risk-based)', 'Checklist used', 'Result', 'Issues found', 'Customer harm? (Yes/No)', 'Action required', 'Corrective Action Log ref.']
    widths = [11, 10, 11, 13, 16, 16, 16, 18, 20, 36, 11, 28, 14]
    ws = b.sheet('File Review Record', widths, 'File Review Record')
    r = b.note(ws, 3, 'One line per file reviewed. Keep the completed checklist and evidence in the Monitoring Evidence Folder [Insert location].', len(cols))
    r = b.header(ws, r, cols, height=40)
    end = b.rows(ws, r, 150, len(cols), height=20, first_ref='FR-')
    b.validate(ws, 'rag', 9, r, end - 1); b.validate(ws, 'yn', 11, r, end - 1)
    b.rag(ws, f'I{r}:I{end - 1}')
    # CAL
    cols = ['CAL ref.', 'Date raised', 'Source', 'Source ref. (e.g. FR-, CB-, CR-)', 'Finding / issue', 'Rating', 'Rule / Principle / policy affected', 'Root cause', 'Action to fix root cause', 'Customers to put right? (details)', 'Owner', 'Target date', 'Status', 'Date completed', 'Verified by (re-test)', 'Date verified closed']
    widths = [10, 11, 18, 14, 34, 20, 20, 26, 32, 22, 13, 11, 16, 11, 14, 11]
    ws = b.sheet('Corrective Action Log', widths, 'Corrective Action Log')
    r = b.note(ws, 3, 'Every Red or Amber finding, breach or significant complaint theme is tracked here to verified closure. Suggested deadlines: Red within [Insert, e.g. 30] days; Amber within [Insert, e.g. 90] days. Overdue actions turn red.', len(cols))
    r = b.header(ws, r, cols, height=40)
    end = b.rows(ws, r, 100, len(cols), height=22, first_ref='CAL-')
    b.validate(ws, 'src', 3, r, end - 1); b.validate(ws, 'rag', 6, r, end - 1); b.validate(ws, 'status', 13, r, end - 1)
    b.rag(ws, f'F{r}:F{end - 1}')
    ws.conditional_formatting.add(f'L{r}:L{end - 1}', FormulaRule(formula=[f'AND(L{r}<>"",P{r}="",L{r}<TODAY())'], fill=PatternFill('solid', fgColor='F4B6B6')))
    # Compliance universe
    cols = ['Ref.', 'Requirement (sourcebook / law)', 'Applies to our activities?', 'Regulated activity', 'Summary of what it requires of us', 'Policy / procedure that covers it', 'Monitoring Plan ref.', 'Owner', 'Last reviewed', 'Regulatory change pending? (details)']
    widths = [7, 34, 11, 20, 40, 26, 12, 14, 11, 28]
    ws = b.sheet('Compliance Universe Register', widths, 'Compliance Universe Register')
    r = b.note(ws, 3, 'List every requirement that applies to your firm. Universal lines are pre-filled; add the activity-specific sourcebooks from section 3 of your Compliance Monitoring Programme under the green OPTION rows.', len(cols))
    r = b.header(ws, r, cols, height=40)
    first = r
    n = 1
    for req, pol in [('Threshold Conditions (COND)', 'Business plan; Risk Management Framework'), ('Principles for Businesses (PRIN)', 'All policies'), ('Consumer Duty (PRIN 2A)', 'Consumer Duty Implementation Plan'), ('SYSC – governance, systems and controls, record keeping, outsourcing', 'Compliance Monitoring Programme; Risk Management Framework'), ('SM&CR, Conduct Rules (COCON), FIT', 'Statements of Responsibilities; Training and Competence Policy'), ('Training and competence (SYSC 5.1.1R; TC where applicable)', 'Training and Competence Policy'), ('Complaints (DISP)', 'Complaints Procedure'), ('Financial promotions (FSMA s21 and activity rules)', 'Financial Promotions Policy'), ('Supervision – notifications and reporting (SUP)', 'Compliance Monitoring Plan'), ('Financial crime (POCA, MLRs where applicable, sanctions, Bribery Act, Criminal Finances Act)', 'AML Policy; CDD Policy'), ('Data protection (UK GDPR, DPA 2018, PECR)', 'Data Processing Policy'), ('Customer vulnerability (FG21/1; Equality Act 2010; Mental Capacity Act 2005)', 'Customer Vulnerability Policy')]:
        r = b.data_row(ws, r, [f'U{n:02d}', req, 'Yes', 'All activities', '[Insert]', pol, '[Insert]', '[Insert]'], len(cols), height=30)
        n += 1
    for act, keep in ACTIVITIES:
        r = b.option_row(ws, r, f'{act} firms', keep, len(cols))
        r = b.data_row(ws, r, [f'U{n:02d}', '[Insert sourcebooks for this activity – see Compliance Monitoring Programme section 3]', 'Yes', act, '[Insert]', '[Insert]', '[Insert]', '[Insert]'], len(cols), height=30)
        n += 1
    b.validate(ws, 'applies', 3, first, r - 1)
    b.save(f'{OUT}/Compliance_Monitoring_Tracker_v2.0_Template.xlsx')


# =====================================================================
# 11. Consumer Duty Evidence (new)
# =====================================================================
def consumer_duty():
    b = Book(BASE_APTOS, 'Aptos Narrow')
    b.guide('How to use', 'Consumer Duty Evidence – how to use', STANDARD_GUIDE + [
        ('h', 'What this workbook contains'),
        ('p', '1. Outcome MI – the management information you use to monitor the four Consumer Duty outcomes (section 8 of the Consumer Duty Implementation Plan), with outcomes for customers with characteristics of vulnerability shown separately. It is the evidence base for your annual board report (PRIN 2A.8).'),
        ('p', '2. Target Market Register – the target market, role and distribution strategy for each product or service (products and services outcome).'),
        ('p', '3. Fair Value Register – each fair value assessment and its conclusion (price and value outcome).'),
        ('p', '4. Board Report Actions – actions agreed from MI reviews and the board report, tracked to completion.'),
        ('p', 'The FCA’s April 2026 observations on board reports stress drawing conclusions from MI (not just presenting it), recording challenge, and covering outcomes through distribution chains and for vulnerable customers.'),
        ('watch', 'CP26/23 (June 2026) proposes clarifying distribution chain responsibilities and board reporting proportionality. Final rules are expected in Q1 2027.'),
    ])
    b.add_list('outcome', ['Products and services', 'Price and value', 'Consumer understanding', 'Consumer support', 'Cross-cutting / all'])
    b.add_list('rag', ['Green – good outcomes', 'Amber – monitor / act', 'Red – poor outcomes'])
    b.add_list('role', ['Manufacturer', 'Co-manufacturer', 'Distributor', 'Manufacturer (own service) and distributor'])
    b.add_list('fv', ['Fair value', 'Fair value – actions required', 'Not fair value – remediation required'])
    b.add_list('status', ['Open', 'In progress', 'Complete'])
    b.add_list('act', ACT_NAMES)
    cols = ['Outcome', 'Metric', 'Definition and data source', 'Tolerance / trigger for action', 'Q1 – all customers', 'Q1 – vulnerable customers', 'Q2 – all customers', 'Q2 – vulnerable customers', 'Q3 – all customers', 'Q3 – vulnerable customers', 'Q4 – all customers', 'Q4 – vulnerable customers', 'RAG', 'Conclusion drawn and action (Board Report Actions ref.)']
    widths = [18, 26, 30, 18, 10, 10, 10, 10, 10, 10, 10, 10, 18, 36]
    ws = b.sheet('Outcome MI', widths, 'Consumer Duty Outcome MI', 'Firm: [Insert firm legal name]   |   Year: [Insert]   |   Owner: [Insert name]')
    r = b.note(ws, 3, 'Starter metrics are suggestions – keep those that are meaningful for your business, add your own, and set a tolerance for each. Compare outcomes for customers with characteristics of vulnerability against all customers.', len(cols))
    r = b.header(ws, r, cols, height=45)
    first = r
    for o, m, d in [('Products and services', 'Sales outside target market', 'Number / % of sales to customers outside the target market (file reviews)'), ('Products and services', 'Early cancellations / exits', '% of customers cancelling or exiting within [Insert period]'), ('Price and value', 'Total cost paid', 'Average fees / charges / commission per customer'), ('Price and value', 'Charges by customer group', 'Default / arrears / additional charges paid by group'), ('Consumer understanding', 'Communication testing', 'Results of comprehension testing of key communications'), ('Consumer understanding', 'Queries showing misunderstanding', 'Number of calls / emails about unclear information'), ('Consumer support', 'Response times', 'Average time to answer / respond by channel'), ('Consumer support', 'Complaints', 'Volume, uphold rate, themes (Complaints Register)'), ('Consumer support', 'Ease of cancelling / switching', 'Time and steps needed to cancel or switch')]:
        r = b.data_row(ws, r, [o, m, d, '[Insert]'], len(cols), height=32)
    end = b.rows(ws, r, 10, len(cols), height=22)
    b.validate(ws, 'outcome', 1, first, end - 1); b.validate(ws, 'rag', 13, first, end - 1)
    b.rag(ws, f'M{first}:M{end - 1}')
    cols = ['Product / service', 'Regulated activity', 'Our role', 'Manufacturer (if not us)', 'Target market (needs, characteristics, objectives)', 'Customers it is NOT suitable for', 'Characteristics of vulnerability considered', 'Distribution strategy / channels', 'Date approved', 'Approved by', 'Last reviewed', 'Next review']
    widths = [22, 20, 18, 18, 36, 28, 26, 24, 11, 14, 11, 11]
    ws = b.sheet('Target Market Register', widths, 'Target Market Register')
    r = b.note(ws, 3, 'One line per product or service you manufacture or distribute. Distributors should record the manufacturer’s target market information and how they sell within it.', len(cols))
    r = b.header(ws, r, cols, height=45)
    end = b.rows(ws, r, 20, len(cols), height=30)
    b.validate(ws, 'act', 2, r, end - 1); b.validate(ws, 'role', 3, r, end - 1)
    cols = ['Assessment ref.', 'Product / service', 'Date of assessment', 'Total price to customer (incl. fees, charges, commission)', 'Benefits to customer', 'Costs to provide', 'Non-financial costs (time, effort, barriers)', 'Groups receiving worse value (incl. vulnerable customers)', 'Distribution chain remuneration considered', 'Conclusion', 'Actions (Board Report Actions ref.)', 'Approved by', 'Next review']
    widths = [11, 22, 11, 26, 26, 18, 22, 24, 22, 20, 22, 14, 11]
    ws = b.sheet('Fair Value Register', widths, 'Fair Value Register')
    r = b.note(ws, 3, 'Record a fair value assessment for each product or service you manufacture (and the manufacturer’s assessment for products you distribute). Review at least annually.', len(cols))
    r = b.header(ws, r, cols, height=55)
    end = b.rows(ws, r, 20, len(cols), height=30, first_ref='FV-')
    b.validate(ws, 'fv', 10, r, end - 1)
    cols = ['Action ref.', 'Date agreed', 'Source (MI review / board report / complaint theme)', 'Outcome affected', 'Issue identified', 'Action', 'Owner', 'Target date', 'Status', 'Evidence of completion', 'Date closed']
    widths = [10, 11, 22, 18, 32, 32, 14, 11, 12, 26, 11]
    ws = b.sheet('Board Report Actions', widths, 'Consumer Duty Board Report Actions')
    r = b.note(ws, 3, 'Track every action arising from outcome monitoring and the annual board report. The board report should show progress against these actions.', len(cols))
    r = b.header(ws, r, cols, height=40)
    end = b.rows(ws, r, 40, len(cols), height=22, first_ref='CD-')
    b.validate(ws, 'outcome', 4, r, end - 1); b.validate(ws, 'status', 9, r, end - 1)
    b.save(f'{OUT}/Consumer_Duty_Evidence_v2.0_Template.xlsx')


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    for fn in [breach_log, cdd_workbook, risk_log, training_log, complaints_register, finprom_log, data_registers, conflicts_gifts, bwra, monitoring_tracker, consumer_duty]:
        fn()
        print('built', fn.__name__)
